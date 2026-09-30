from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.importers.idms_client import IdmsClient, MfaRequired
from app.importers.idms_parsers import (
    IDMS_CHARGE_OFF_REPORT_ID,
    IDMS_MONTH_END_REPORT_ID,
    IDMS_SALES_REPORT_ID,
    parse_aa_month_end,
    parse_aa_sales_manual,
    parse_inventory,
    parse_report,
)
from app.models.idms_inventory import IdmsInventory
from app.repositories.idms_repo import IdmsRepository
from app.schemas.idms import IdmsSyncOut
from app.services.idms_inventory_pdf import build_inventory_pdf

INVENTORY_XLSX_HEADERS = [
    "Stock #", "Year", "Make", "Model/Trim", "Color", "VIN Last 6",
    "Mileage", "Status", "Alternate Lot", "Acquired Date", "DOL",
    "Asking Price", "Wholesale $", "Inventory Flags",
]
_INVENTORY_MONEY_COLS = (12, 13)  # Asking Price, Wholesale $ (1-indexed)
_INVENTORY_DATE_COL = 10  # Acquired Date


def _vehicle_label(row: IdmsInventory) -> str:
    desc = " ".join(p for p in (row.vehicle_year, row.make, row.model_trim) if p)
    return f"{row.stock_number or '—'} — {desc}" if desc else (row.stock_number or "—")


def _channel_breakdown(rows: list[IdmsInventory]) -> list[tuple[str, list[str]]]:
    """Agrupa por canal (split de inventory_flags por coma) las descripciones
    de vehículo ('Stock # — Año Marca Modelo/Trim'), ordenado por # unidades."""
    grouped: dict[str, list[str]] = {}
    for row in rows:
        if not row.inventory_flags:
            continue
        label = _vehicle_label(row)
        for flag in row.inventory_flags.split(","):
            flag = flag.strip()
            if flag:
                grouped.setdefault(flag, []).append(label)
    return sorted(grouped.items(), key=lambda kv: len(kv[1]), reverse=True)


def _write_dashboard_sheet(ws, rows: list[IdmsInventory], kpis: dict, aging: list[dict]) -> None:
    title_font = Font(bold=True, size=16)
    label_font = Font(bold=True, size=9, color="666666")
    value_font = Font(bold=True, size=18)
    header_font = Font(bold=True)
    section_font = Font(bold=True, size=12)

    ws["A1"] = "Reporte de Inventario — Automania"
    ws["A1"].font = title_font
    ws.merge_cells("A1:F1")

    snapshot = kpis.get("snapshot_date")
    ws["A2"] = f"Snapshot: {snapshot}" if snapshot else "Snapshot: —"
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells("A2:F2")

    kpi_cards = [
        ("UNIDADES", kpis["count"], "#,##0"),
        ("ASKING PRICE TOTAL", float(kpis["total_price"]), "$#,##0.00"),
        ("WHOLESALE TOTAL", float(kpis["total_wholesale"]), "$#,##0.00"),
        ("DOL PROMEDIO", round(kpis["avg_dol"], 1), "0.0"),
        ("DOL MÁXIMO", kpis["max_dol"], "#,##0"),
    ]
    for (label, value, fmt), (c1, c2) in zip(kpi_cards, [("A", "B"), ("C", "D"), ("E", "F"), ("G", "H"), ("I", "J")]):
        ws.merge_cells(f"{c1}4:{c2}4")
        ws[f"{c1}4"] = label
        ws[f"{c1}4"].font = label_font
        ws.merge_cells(f"{c1}5:{c2}5")
        cell = ws[f"{c1}5"]
        cell.value = value
        cell.font = value_font
        cell.number_format = fmt

    # ------------------------------------------------------------------
    # Antigüedad en lote (DOL)
    # ------------------------------------------------------------------
    aging_row = 8
    ws.cell(row=aging_row, column=1, value="Antigüedad en lote (DOL)").font = section_font
    ws.cell(row=aging_row + 1, column=1, value="Bucket").font = header_font
    ws.cell(row=aging_row + 1, column=2, value="Unidades").font = header_font
    ws.cell(row=aging_row + 1, column=3, value="Asking Price").font = header_font
    for i, bucket in enumerate(aging):
        r = aging_row + 2 + i
        ws.cell(row=r, column=1, value=bucket["bucket"])
        ws.cell(row=r, column=2, value=bucket["count"])
        price_cell = ws.cell(row=r, column=3, value=float(bucket["total_price"]))
        price_cell.number_format = "$#,##0.00"
    aging_last_row = aging_row + 1 + max(len(aging), 1)

    if aging:
        aging_chart = BarChart()
        aging_chart.title = "Unidades por antigüedad (DOL)"
        aging_chart.y_axis.title = "Unidades"
        aging_chart.style = 10
        aging_chart.width = 14
        aging_chart.height = 8
        data = Reference(ws, min_col=2, min_row=aging_row + 1, max_row=aging_last_row)
        cats = Reference(ws, min_col=1, min_row=aging_row + 2, max_row=aging_last_row)
        aging_chart.add_data(data, titles_from_data=True)
        aging_chart.set_categories(cats)
        ws.add_chart(aging_chart, f"E{aging_row}")

    # ------------------------------------------------------------------
    # Publicado en (canales)
    # ------------------------------------------------------------------
    channels = _channel_breakdown(rows)
    ch_row = aging_last_row + 3
    ws.cell(row=ch_row, column=1, value="Publicado en (canales)").font = section_font
    ws.cell(row=ch_row + 1, column=1, value="Canal").font = header_font
    ws.cell(row=ch_row + 1, column=2, value="Unidades").font = header_font
    ws.cell(row=ch_row + 1, column=12, value="Vehículos").font = header_font  # col L
    wrap_top = Alignment(wrap_text=True, vertical="top")
    for i, (channel, labels) in enumerate(channels):
        r = ch_row + 2 + i
        ws.cell(row=r, column=1, value=channel).alignment = wrap_top
        ws.cell(row=r, column=2, value=len(labels)).alignment = wrap_top
        vehicles_cell = ws.cell(row=r, column=12, value="\n".join(labels))  # col L
        vehicles_cell.alignment = wrap_top
        ws.row_dimensions[r].height = max(15, 14 * len(labels))
    ch_last_row = ch_row + 1 + max(len(channels), 1)

    if channels:
        ch_chart = BarChart()
        ch_chart.title = "Unidades por canal"
        ch_chart.y_axis.title = "Unidades"
        ch_chart.style = 11
        ch_chart.width = 14
        ch_chart.height = 8
        data = Reference(ws, min_col=2, min_row=ch_row + 1, max_row=ch_last_row)
        cats = Reference(ws, min_col=1, min_row=ch_row + 2, max_row=ch_last_row)
        ch_chart.add_data(data, titles_from_data=True)
        ch_chart.set_categories(cats)
        # Columna N, lejos de E (chart de aging) y de L (lista de vehículos,
        # columna ancha) — evita que los dos charts se superpongan aunque sus
        # filas se crucen verticalmente (el de aging puede medir ~15 filas).
        ws.add_chart(ch_chart, f"N{ch_row}")

    for col, width in zip("ABCD", [22, 14, 16, 10]):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["L"].width = 60


def _write_detail_sheet(ws, rows: list[IdmsInventory]) -> None:
    ws.append(INVENTORY_XLSX_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([
            row.stock_number,
            row.vehicle_year,
            row.make,
            row.model_trim,
            row.exterior_color,
            row.vin_last6,
            row.mileage,
            row.status,
            row.alternate_lot,
            row.acq_date,
            row.dol,
            float(row.price),
            float(row.wholesale_price),
            row.inventory_flags,
        ])

    for excel_row in ws.iter_rows(min_row=2):
        for col_idx in _INVENTORY_MONEY_COLS:
            excel_row[col_idx - 1].number_format = "$#,##0.00"
        excel_row[_INVENTORY_DATE_COL - 1].number_format = "mm/dd/yyyy"

    for i, header in enumerate(INVENTORY_XLSX_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(header) + 2, 12)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_inventory_xlsx(rows: list[IdmsInventory], kpis: dict, aging: list[dict]) -> bytes:
    """Genera el .xlsx de inventario: hoja Dashboard (KPIs + gráficas) e
    Inventario (detalle con AutoFilter para filtrar en Excel)."""
    wb = Workbook()
    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"
    _write_dashboard_sheet(ws_dashboard, rows, kpis, aging)

    ws_detail = wb.create_sheet("Inventario")
    _write_detail_sheet(ws_detail, rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class IdmsService:
    def __init__(self, repo: IdmsRepository | None = None):
        self.repo = repo or IdmsRepository()

    # ------------------------------------------------------------------
    # Session / login
    # ------------------------------------------------------------------
    def check_session(self) -> dict:
        client = IdmsClient()
        if client.load_session() and client.is_authenticated():
            return {"authenticated": True, "message": "Sesión activa"}
        return {"authenticated": False, "message": "Sin sesión activa"}

    def login(self, otp_code: str | None = None) -> dict:
        client = IdmsClient()
        try:
            client.login(otp_code=otp_code)
        except MfaRequired as exc:
            return {"authenticated": False, "mfa_required": True, "message": exc.message}
        return {"authenticated": True, "message": "Sesión iniciada"}

    # ------------------------------------------------------------------
    # Charge Offs
    # ------------------------------------------------------------------
    async def sync_charge_offs(
        self, db: AsyncSession, year: int
    ) -> IdmsSyncOut:
        client = IdmsClient()
        client.login()
        raw = client.export_csv(IDMS_CHARGE_OFF_REPORT_ID, export_type="csv")
        rows = parse_report(IDMS_CHARGE_OFF_REPORT_ID, raw)
        # Se sincronizan solo los años presentes en el reporte descargado,
        # preservando los históricos cargados desde Excel.
        inserted = await self.repo.sync_charge_off_years(db, rows)
        return IdmsSyncOut(
            report_id=IDMS_CHARGE_OFF_REPORT_ID,
            year=year,
            rows_inserted=inserted,
            message=f"{inserted} charge offs sincronizados",
        )

    async def import_charge_off_historical(
        self, db: AsyncSession, file_bytes: bytes
    ) -> IdmsSyncOut:
        from app.importers.idms_chargeoff_excel import parse_chargeoff_historical_excel

        rows = parse_chargeoff_historical_excel(file_bytes)
        inserted = await self.repo.sync_all_charge_offs(db, rows)
        return IdmsSyncOut(
            report_id="manual-pull",
            year=0,
            rows_inserted=inserted,
            message=f"{inserted} charge offs históricos importados",
        )

    async def get_charge_offs(self, db: AsyncSession, year: int):
        return await self.repo.list_charge_offs(db, report_year=year)

    async def get_charge_off_kpis(self, db: AsyncSession, year: int):
        return await self.repo.get_charge_off_kpis(db, report_year=year)

    async def get_charge_off_monthly(self, db: AsyncSession, year: int):
        return await self.repo.get_charge_off_monthly(db, report_year=year)

    async def get_available_years(self, db: AsyncSession) -> list[int]:
        return await self.repo.get_available_years(db)

    async def get_charge_off_overview(self, db: AsyncSession, year: int):
        return await self.repo.get_charge_off_overview(db, year)

    async def get_charge_off_monthly_detail(self, db: AsyncSession, year: int):
        return await self.repo.get_charge_off_monthly_detail(db, year)

    # ------------------------------------------------------------------
    # Month End (snapshot de cartera)
    # ------------------------------------------------------------------
    async def sync_month_end(self, db: AsyncSession) -> IdmsSyncOut:
        """Toma la foto de la cartera de hoy y la guarda como el mes corriente."""
        snapshot = date.today()
        client = IdmsClient()
        client.login()
        raw = client.export_csv(IDMS_MONTH_END_REPORT_ID, export_type="csv")
        rows = parse_aa_month_end(raw, snapshot)
        inserted = await self.repo.sync_month_end(
            db, rows, year=snapshot.year, month=snapshot.month
        )
        return IdmsSyncOut(
            report_id=IDMS_MONTH_END_REPORT_ID,
            year=snapshot.year,
            rows_inserted=inserted,
            message=(
                f"{inserted} cuentas en el snapshot de "
                f"{snapshot.strftime('%m/%Y')}"
            ),
        )

    # ------------------------------------------------------------------
    # Sales
    # ------------------------------------------------------------------
    async def sync_sales(
        self, db: AsyncSession, year: int
    ) -> IdmsSyncOut:
        client = IdmsClient()
        client.login()
        raw = client.export_csv(IDMS_SALES_REPORT_ID, export_type="csv")
        rows = parse_report(IDMS_SALES_REPORT_ID, raw)
        inserted = await self.repo.sync_sales(db, rows)
        return IdmsSyncOut(
            report_id=IDMS_SALES_REPORT_ID,
            year=year,
            rows_inserted=inserted,
            message=f"{inserted} ventas sincronizadas",
        )

    async def import_sales_historical(
        self, db: AsyncSession, file_bytes: bytes
    ) -> IdmsSyncOut:
        rows = parse_aa_sales_manual(file_bytes)
        inserted = await self.repo.sync_all_sales(db, rows)
        return IdmsSyncOut(
            report_id="manual-pull",
            year=0,
            rows_inserted=inserted,
            message=f"{inserted} ventas históricas importadas",
        )

    async def get_sales(self, db: AsyncSession, year: int):
        return await self.repo.list_sales(db, report_year=year)

    async def get_sales_kpis(self, db: AsyncSession, year: int):
        return await self.repo.get_sales_kpis(db, report_year=year)

    async def get_sales_monthly(self, db: AsyncSession, year: int):
        return await self.repo.get_sales_monthly(db, report_year=year)

    async def get_sales_years(self, db: AsyncSession) -> list[int]:
        return await self.repo.get_sales_years(db)

    async def get_sales_by_salesperson(self, db: AsyncSession, year: int):
        return await self.repo.get_sales_by_salesperson(db, report_year=year)

    async def get_sales_by_vehicle(self, db: AsyncSession, year: int):
        return await self.repo.get_sales_by_vehicle(db, report_year=year)

    # ------------------------------------------------------------------
    # Inventario
    # ------------------------------------------------------------------
    async def sync_inventory(
        self, db: AsyncSession, status: str = "A"
    ) -> IdmsSyncOut:
        client = IdmsClient()
        client.login()
        html = client.fetch_inventory(status=status)
        rows = parse_inventory(html)
        inserted = await self.repo.sync_inventory(db, rows)
        return IdmsSyncOut(
            report_id="inventory-lookup",
            year=0,
            rows_inserted=inserted,
            message=f"{inserted} unidades de inventario sincronizadas",
        )

    async def get_inventory(self, db: AsyncSession):
        return await self.repo.list_inventory(db)

    async def get_inventory_kpis(self, db: AsyncSession):
        return await self.repo.get_inventory_kpis(db)

    async def get_inventory_aging(self, db: AsyncSession):
        return await self.repo.get_inventory_aging(db)

    async def export_inventory_xlsx(self, db: AsyncSession) -> bytes:
        rows = await self.repo.list_inventory(db)
        kpis = await self.repo.get_inventory_kpis(db)
        aging = await self.repo.get_inventory_aging(db)
        return build_inventory_xlsx(rows, kpis, aging)

    async def export_inventory_pdf(self, db: AsyncSession) -> bytes:
        rows = await self.repo.list_inventory(db)
        kpis = await self.repo.get_inventory_kpis(db)
        aging = await self.repo.get_inventory_aging(db)
        return build_inventory_pdf(rows, kpis, aging)
